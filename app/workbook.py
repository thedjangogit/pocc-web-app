"""
Reads and writes reviewer decisions in the validated workbooks -- the
same files the pipeline reads, so a decision saved here is what the next
pipeline run acts on. Rows are keyed on the sheet's first two columns
(left id, right id): re-deciding a pair overwrites its row instead of
adding a second one.

Every write reloads the workbook from disk immediately before changing
it, so a decision another reviewer saved since this session last read
the file is kept, not overwritten. That narrows, but does not close, the
window for two people saving at the same moment -- acceptable for the
pilot; see README.md.

"Unsure" is not a value the pipeline accepts in is_duplicate, so unsure
items go to a separate append-only CSV next to the workbooks instead.
"""

import csv
import os
from datetime import datetime

import openpyxl
import pandas as pd
import streamlit as st

from config import COMMENT_COL, IS_DUPLICATE_COL, KEEP_COL


class WorkbookBusy(Exception):
    pass


def _lock_files(path: str) -> list[str]:
    folder, name = os.path.split(path)
    # Excel's lock file is "~$" + name; Word-style truncation swaps the
    # first two characters instead. Check both.
    candidates = {os.path.join(folder, "~$" + name), os.path.join(folder, "~$" + name[2:])}
    return [p for p in candidates if os.path.exists(p)]


def _norm(value) -> str:
    return "" if value is None else str(value).strip()


@st.cache_data(show_spinner=False)
def read_decisions(path: str, file_mtime: float) -> dict[str, dict[tuple[str, str], dict]]:
    """{sheet: {(left_id, right_id): {is_duplicate, keep, comment}}}"""
    if not os.path.exists(path):
        return {}
    sheets = pd.read_excel(path, sheet_name=None, dtype=str, keep_default_na=False)
    result = {}
    for sheet, df in sheets.items():
        if df.shape[1] < 2:
            continue
        left_col, right_col = df.columns[0], df.columns[1]
        decisions = {}
        for rec in df.to_dict("records"):
            key = (_norm(rec[left_col]), _norm(rec[right_col]))
            if key == ("", ""):
                continue
            decisions[key] = {
                "is_duplicate": _norm(rec.get(IS_DUPLICATE_COL)).upper(),
                "keep": _norm(rec.get(KEEP_COL)).lower(),
                "comment": _norm(rec.get(COMMENT_COL)),
            }
        result[sheet] = decisions
    return result


def _open_sheet(path: str, sheet: str):
    if not os.path.exists(path):
        raise FileNotFoundError(f"Validated workbook not found: {path}")
    locks = _lock_files(path)
    if locks:
        raise WorkbookBusy(
            f"{os.path.basename(path)} is open in Excel (lock file {os.path.basename(locks[0])}). "
            "Close it there, then click again."
        )
    wb = openpyxl.load_workbook(path)
    if sheet not in wb.sheetnames:
        raise KeyError(f"Sheet '{sheet}' not found in {os.path.basename(path)}")
    ws = wb[sheet]
    headers = [_norm(c.value) for c in ws[1]]
    return wb, ws, headers


def _rows_for_pairs(ws, pairs: set[tuple[str, str]]) -> dict[tuple[str, str], int]:
    found = {}
    for r in range(2, ws.max_row + 1):
        key = (_norm(ws.cell(r, 1).value), _norm(ws.cell(r, 2).value))
        if key in pairs:
            found[key] = r
    return found


def _save_atomic(wb, path: str) -> None:
    folder, name = os.path.split(path)
    tmp = os.path.join(folder, f".{name}.{os.getpid()}.tmp")
    wb.save(tmp)
    os.replace(tmp, path)


def write_decision(path: str, sheet: str, pairs, is_duplicate: str, keep: str, comment: str) -> None:
    wb, ws, headers = _open_sheet(path, sheet)
    col = {h: i + 1 for i, h in enumerate(headers) if h}
    values = {IS_DUPLICATE_COL: is_duplicate, KEEP_COL: keep or None, COMMENT_COL: comment or None}

    existing = _rows_for_pairs(ws, set(pairs))
    for left, right in pairs:
        r = existing.get((left, right))
        if r is None:
            r = ws.max_row + 1
            ws.cell(r, 1, left)
            ws.cell(r, 2, right)
        for header, value in values.items():
            ws.cell(r, col[header], value)
    _save_atomic(wb, path)


def clear_decision(path: str, sheet: str, pairs) -> None:
    wb, ws, _ = _open_sheet(path, sheet)
    rows = sorted(_rows_for_pairs(ws, set(pairs)).values(), reverse=True)
    if not rows:
        return
    for r in rows:
        ws.delete_rows(r)
    _save_atomic(wb, path)


UNSURE_FIELDS = ["timestamp", "reviewer", "action", "item_key", "workbook", "sheet", "pairs", "comment"]


@st.cache_data(show_spinner=False)
def read_unsure(path: str, file_mtime: float) -> dict[str, dict]:
    """Latest log entry per item; items whose latest action is 'cleared' are dropped."""
    if not os.path.exists(path):
        return {}
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    latest = {rec["item_key"]: rec for rec in df.to_dict("records")}
    return {k: v for k, v in latest.items() if v["action"] == "unsure"}


def log_unsure(path: str, action: str, item_key: str, workbook: str, sheet: str, pairs, reviewer: str, comment: str) -> None:
    new_file = not os.path.exists(path)
    with open(path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=UNSURE_FIELDS)
        if new_file:
            writer.writeheader()
        writer.writerow(
            {
                "timestamp": datetime.now().isoformat(timespec="seconds"),
                "reviewer": reviewer,
                "action": action,
                "item_key": item_key,
                "workbook": workbook,
                "sheet": sheet,
                "pairs": "; ".join(f"{l} | {r}" for l, r in pairs),
                "comment": comment,
            }
        )
